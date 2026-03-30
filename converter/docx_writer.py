"""
DOCX 文档生成模块
接收 mistune AST 中间结构，根据样式模式生成格式化的 Word 文档。
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

from styles.docx_styles import (
    get_heading_map,
    FONT_CONFIG,
    LINE_SPACING,
    SPACE_BEFORE,
    SPACE_AFTER,
    HEADING_FONT_SIZES,
    CODE_BLOCK_BG_COLOR,
)
from converter.md_parser import extract_text_from_children


class DocxWriter:
    """Word 文档生成器"""

    def __init__(self, style_mode: str = 'A'):
        """
        Args:
            style_mode: 'A'（标题为Title）或 'B'（标题为Heading1）
        """
        self.style_mode = style_mode
        self.heading_map = get_heading_map(style_mode)
        self.doc = Document()
        self._setup_default_styles()

    def _setup_default_styles(self):
        """设置文档默认样式。"""
        style = self.doc.styles['Normal']
        font = style.font
        font.name = FONT_CONFIG['body']['name_en']
        font.size = FONT_CONFIG['body']['size']
        style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CONFIG['body']['name_cn'])

        # 设置默认段落格式
        pf = style.paragraph_format
        pf.line_spacing = LINE_SPACING
        pf.space_before = SPACE_BEFORE
        pf.space_after = SPACE_AFTER

    def convert(self, ast_nodes: list, output_path: str):
        """将 AST 节点列表转换为 DOCX 并保存。

        Args:
            ast_nodes: mistune AST 节点列表
            output_path: 输出文件路径
        """
        for node in ast_nodes:
            self._process_node(node)

        # 确保输出目录存在
        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self.doc.save(output_path)

    def _process_node(self, node: dict, list_level: int = 0):
        """处理单个 AST 节点。

        Args:
            node: AST 节点字典
            list_level: 列表嵌套层级
        """
        node_type = node.get('type', '')

        if node_type == 'heading':
            self._add_heading(node)
        elif node_type == 'paragraph':
            self._add_paragraph(node)
        elif node_type == 'list':
            self._add_list(node, list_level)
        elif node_type == 'block_code':
            self._add_code_block(node)
        elif node_type == 'table':
            self._add_table(node)
        elif node_type == 'thematic_break':
            self._add_horizontal_rule()
        elif node_type == 'block_quote':
            self._add_block_quote(node)
        elif node_type == 'blank_line':
            pass  # 忽略空行
        else:
            # 对于未知类型，尝试提取文本
            if 'children' in node:
                for child in node['children']:
                    if isinstance(child, dict):
                        self._process_node(child, list_level)

    def _add_heading(self, node: dict):
        """添加标题。"""
        level = node.get('attrs', {}).get('level', 1)
        children = node.get('children', [])
        text = extract_text_from_children(children)

        style_name = self.heading_map.get(level, f'Heading {min(level, 9)}')

        paragraph = self.doc.add_paragraph(style=style_name)
        self._add_inline_content(paragraph, children)

        # 设置标题字体
        font_size = HEADING_FONT_SIZES.get(style_name)
        if font_size:
            for run in paragraph.runs:
                run.font.size = font_size
                run.font.name = FONT_CONFIG['heading']['name_en']
                run.font.element.rPr.rFonts.set(
                    qn('w:eastAsia'), FONT_CONFIG['heading']['name_cn']
                )

    def _add_paragraph(self, node: dict):
        """添加段落。"""
        children = node.get('children', [])
        paragraph = self.doc.add_paragraph()
        self._add_inline_content(paragraph, children)

    def _add_inline_content(self, paragraph, children: list):
        """向段落中添加行内内容（文本、加粗、斜体、代码等）。"""
        if not children:
            return

        for child in children:
            if isinstance(child, str):
                run = paragraph.add_run(child)
                self._apply_body_font(run)
            elif isinstance(child, dict):
                child_type = child.get('type', '')

                if child_type == 'text':
                    raw = child.get('raw', '') or child.get('text', '')
                    run = paragraph.add_run(raw)
                    self._apply_body_font(run)

                elif child_type == 'strong':
                    inner_children = child.get('children', [])
                    text = extract_text_from_children(inner_children)
                    run = paragraph.add_run(text)
                    run.bold = True
                    self._apply_body_font(run)

                elif child_type == 'emphasis':
                    inner_children = child.get('children', [])
                    text = extract_text_from_children(inner_children)
                    run = paragraph.add_run(text)
                    run.italic = True
                    self._apply_body_font(run)

                elif child_type == 'codespan':
                    raw = child.get('raw', '') or child.get('text', '')
                    run = paragraph.add_run(raw)
                    run.font.name = FONT_CONFIG['code']['name_en']
                    run.font.size = FONT_CONFIG['code']['size']
                    # 给行内代码加灰色背景
                    shd = parse_xml(
                        f'<w:shd {nsdecls("w")} w:fill="F0F0F0" w:val="clear"/>'
                    )
                    run.element.rPr.append(shd)

                elif child_type == 'link':
                    inner_children = child.get('children', [])
                    text = extract_text_from_children(inner_children)
                    link = child.get('link', '') or child.get('attrs', {}).get('url', '')
                    run = paragraph.add_run(text)
                    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
                    run.underline = True
                    self._apply_body_font(run)

                elif child_type == 'softbreak':
                    run = paragraph.add_run('\n')

                elif child_type == 'linebreak':
                    run = paragraph.add_run('\n')

                elif child_type == 'image':
                    alt = child.get('alt', '') or child.get('attrs', {}).get('alt', '')
                    src = child.get('src', '') or child.get('attrs', {}).get('url', '')
                    # 图片引入——如果是本地路径且存在
                    if src and os.path.isfile(src):
                        try:
                            paragraph.add_run().add_picture(src, width=Inches(5))
                        except Exception:
                            run = paragraph.add_run(f'[图片: {alt or src}]')
                            self._apply_body_font(run)
                    else:
                        run = paragraph.add_run(f'[图片: {alt or src}]')
                        self._apply_body_font(run)

                elif 'children' in child:
                    self._add_inline_content(paragraph, child['children'])

    def _add_list(self, node: dict, level: int = 0):
        """添加列表（有序/无序）。"""
        ordered = node.get('attrs', {}).get('ordered', False)
        children = node.get('children', [])
        counter = node.get('attrs', {}).get('start', 1) or 1

        for item in children:
            if not isinstance(item, dict):
                continue
            item_type = item.get('type', '')
            if item_type == 'list_item':
                item_children = item.get('children', [])
                for sub_node in item_children:
                    if not isinstance(sub_node, dict):
                        continue
                    if sub_node.get('type') == 'paragraph':
                        # 生成列表前缀
                        if ordered:
                            prefix = f'{counter}. '
                            counter += 1
                        else:
                            prefix = '• '

                        indent = '    ' * level
                        paragraph = self.doc.add_paragraph()
                        # 设置缩进
                        paragraph.paragraph_format.left_indent = Cm(1.27 * (level + 1))

                        run = paragraph.add_run(indent + prefix)
                        self._apply_body_font(run)
                        self._add_inline_content(paragraph, sub_node.get('children', []))

                    elif sub_node.get('type') == 'list':
                        # 嵌套列表
                        self._add_list(sub_node, level + 1)

    def _add_code_block(self, node: dict):
        """添加代码块。"""
        raw = node.get('raw', '') or node.get('text', '')
        # 移除末尾多余的换行
        raw = raw.rstrip('\n')

        paragraph = self.doc.add_paragraph()

        # 设置代码块背景
        pPr = paragraph._element.get_or_add_pPr()
        shd = parse_xml(
            f'<w:shd {nsdecls("w")} w:fill="F5F5F5" w:val="clear"/>'
        )
        pPr.append(shd)

        # 设置边框
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="4" w:color="DDDDDD"/>'
            f'  <w:left w:val="single" w:sz="4" w:space="4" w:color="DDDDDD"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="4" w:color="DDDDDD"/>'
            f'  <w:right w:val="single" w:sz="4" w:space="4" w:color="DDDDDD"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

        # 设置段落间距
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)

        run = paragraph.add_run(raw)
        run.font.name = FONT_CONFIG['code']['name_en']
        run.font.size = FONT_CONFIG['code']['size']

    def _add_table(self, node: dict):
        """添加表格。"""
        children = node.get('children', [])
        if not children:
            return

        # 收集所有行数据：每行是一个 cell 列表
        all_rows = []
        head_row_count = 0

        for child in children:
            if not isinstance(child, dict):
                continue
            if child.get('type') == 'table_head':
                # table_head 直接包含 table_cell（它本身就是表头行）
                cells = child.get('children', [])
                if cells:
                    all_rows.append(cells)
                    head_row_count = 1
            elif child.get('type') == 'table_body':
                # table_body 包含 table_row，每个 table_row 包含 table_cell
                for row_node in child.get('children', []):
                    if isinstance(row_node, dict) and row_node.get('type') == 'table_row':
                        all_rows.append(row_node.get('children', []))

        if not all_rows:
            return

        # 计算列数
        cols = max(len(row) for row in all_rows) if all_rows else 0
        if cols == 0:
            return

        # 创建表格
        table = self.doc.add_table(rows=len(all_rows), cols=cols, style='Table Grid')

        for row_idx, row_cells in enumerate(all_rows):
            for col_idx, cell_node in enumerate(row_cells):
                if col_idx < cols and isinstance(cell_node, dict):
                    cell = table.cell(row_idx, col_idx)
                    cell_text = extract_text_from_children(cell_node.get('children', []))
                    cell.text = cell_text

                    # 表头加粗
                    if row_idx < head_row_count:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True

    def _add_horizontal_rule(self):
        """添加分隔线。"""
        paragraph = self.doc.add_paragraph()
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'  <w:bottom w:val="single" w:sz="12" w:space="1" w:color="AAAAAA"/>'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

    def _add_block_quote(self, node: dict):
        """添加引用块。"""
        children = node.get('children', [])
        for child in children:
            if isinstance(child, dict):
                if child.get('type') == 'paragraph':
                    paragraph = self.doc.add_paragraph()
                    # 添加左侧蓝色边框
                    pPr = paragraph._element.get_or_add_pPr()
                    pBdr = parse_xml(
                        f'<w:pBdr {nsdecls("w")}>'
                        f'  <w:left w:val="single" w:sz="18" w:space="8" w:color="4A90D9"/>'
                        f'</w:pBdr>'
                    )
                    pPr.append(pBdr)

                    # 设置灰色背景
                    shd = parse_xml(
                        f'<w:shd {nsdecls("w")} w:fill="F0F5FA" w:val="clear"/>'
                    )
                    pPr.append(shd)

                    # 设置缩进
                    paragraph.paragraph_format.left_indent = Cm(1)

                    self._add_inline_content(paragraph, child.get('children', []))
                else:
                    self._process_node(child)

    def _apply_body_font(self, run):
        """给 run 应用正文字体。"""
        run.font.name = FONT_CONFIG['body']['name_en']
        run.font.size = FONT_CONFIG['body']['size']
        run.font.element.rPr.rFonts.set(
            qn('w:eastAsia'), FONT_CONFIG['body']['name_cn']
        )
